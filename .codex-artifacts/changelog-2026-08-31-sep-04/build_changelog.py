from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from lxml import etree


REFERENCE = Path(r"C:\Users\young\Downloads\15 - 19 June.docx")
OUTPUT = Path(r"D:\Arclight\Arclight_App_NewApp_ChatRefact\Arclight_App\31 August - 4 September.docx")
TASK_DIR = Path(r"D:\Arclight\Arclight_App_NewApp_ChatRefact\Arclight_App\.codex-artifacts\changelog-2026-08-31-sep-04")
EXPECTED_SHA256 = "0B19F0022172001957D1D9703465646C89E02A3AECA2DB284FFAE95386EC5122"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
DC_NS = "http://purl.org/dc/elements/1.1/"
DCTERMS_NS = "http://purl.org/dc/terms/1.1/"
NS = {"w": W_NS, "dc": DC_NS, "dcterms": DCTERMS_NS}

BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
TABLE_WIDTH_DXA = 15240
COLUMN_WIDTHS_DXA = [2160, 3420, 4800, 4860]

SECTIONS: list[tuple[str, list[str]]] = [
    (
        "Narration and Video Learning",
        [
            "Added English, Latin American Spanish and Korean narration with matching subtitles to the Fundal Reflex full animation.",
            "Added an accessible menu to choose the narration language, follow the app language or turn narration off.",
            "Replaced the Fundal Reflex full-animation video and improved its timing so the narration, subtitles and animation stay in step.",
            "Added the same multilingual narration to the combined Fundal Reflex examination scroll lesson.",
            "Added English narration and subtitles to the Direct Ophthalmoscopy full animation and replaced its video with an updated version.",
            "Improved iPhone playback to reduce narration drop-outs while large videos load.",
        ],
    ),
    (
        "Language and Offline Access",
        [
            "Expanded Spanish translations across the main app and mini apps.",
            "Included narration audio and subtitle files in offline downloads, with the available language matched to the user's choice.",
            "Improved narration and subtitle controls on smaller screens.",
        ],
    ),
    (
        "Reliability and Testing",
        [
            "Added automatic checks for narration files, subtitle timing, offline downloads, iPhone playback and scroll narration.",
            "Updated a web request package to a newer version to improve security.",
        ],
    ),
    (
        "Work in Progress",
        [
            "Narration and translated subtitle work for the Binocular Indirect Ophthalmoscopy (BIO) and Front of Eye full animations is in progress.",
        ],
    ),
]

TABLE_HEADERS = ["Area", "Before", "Now", "Practical benefit"]
TABLE_ROWS = [
    [
        "Total scheduled runs",
        "About 20 runs per week",
        "Six runs per week",
        "Uses less allowance and focuses on essential checks",
    ],
    [
        "Operations monitoring",
        "Daily at 04:00 and weekdays at 08:30",
        "Monday, Wednesday and Friday at 08:30",
        "Removes repetitive status checks",
    ],
    [
        "Detecting new updates",
        "Mainly checked recent activity on main",
        "Checks changes since the previous run and distinguishes pushes to other branches",
        "Less likely to miss recently pushed work",
    ],
    [
        "Product and UI review",
        "Wednesday and translation checks ran separately",
        "Combined into one Wednesday 04:00 review",
        "UI, accessibility, screen sizes and language findings arrive together",
    ],
    [
        "Full translation review",
        "Split across four monthly runs",
        "Included in the first Wednesday report each month",
        "Retains full monthly coverage with fewer runs",
    ],
    [
        "Security and housekeeping",
        "Security ran on Friday while repository hygiene ran on Monday",
        "Combined into Friday's Security + Hygiene review",
        "Security, refactoring and unnecessary-file candidates appear in one report",
    ],
    [
        "Usage control",
        "Work could stop when a 10% monthly allowance could not be calculated",
        "Scope adjusts gradually using actual rolling seven-day usage",
        "Prevents jobs being skipped despite available capacity",
    ],
    [
        "Report format",
        "Most results were brief chat messages",
        "Major reviews provide detailed chat reports and a DOCX when supported",
        "Results are easier to understand and share",
    ],
    [
        "PostgreSQL DOCX",
        "Could not be created without the reporting connection",
        "Ready for private server-side generation once connected",
        "It will remain unavailable until Arclight Reporting is connected",
    ],
    [
        "Safety",
        "Read-only guidance existed",
        "Explicit read-only restrictions now apply to every active job",
        "Scheduled work cannot change code, deployments or database records",
    ],
    [
        "Agent names",
        "Run days were not obvious",
        "[Mon/Wed/Fri]Operations, [Wed]Product + Language, [Fri]Security + Hygiene, [Fri]Weekly Reports",
        "Each name clearly shows when the agent runs and what it does",
    ],
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def set_run_font(run, *, size: float, bold: bool = False, colour=BLACK) -> None:
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour
    run.font.underline = False


def set_spacing(paragraph, *, before: float = 0, after: float = 0, line: float = 1.05) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def remove_paragraph_border(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.keep_with_next = True
    set_spacing(paragraph, after=4, line=1.0)
    remove_paragraph_border(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=14, bold=True)


def add_intro(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    set_spacing(paragraph, after=8, line=1.08)
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    set_spacing(paragraph, before=8, after=2, line=1.0)
    remove_paragraph_border(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, bold=True)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Paragraph")
    set_spacing(paragraph, after=1, line=1.05)
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.append(numpr)
    for child in list(numpr):
        numpr.remove(child)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), "36")
    numpr.extend([ilvl, numid])
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def clear_document_body(document: Document) -> None:
    body = document._body._element
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_fill(cell, colour: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), colour)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "D9D9D9")

    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, COLUMN_WIDTHS_DXA, strict=True):
        grid_col.set(qn("w:w"), str(width))

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell, width in zip(row.cells, COLUMN_WIDTHS_DXA, strict=True):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def fill_cell(cell, text: str, *, header: bool = False, label: bool = False, shaded: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_spacing(paragraph, line=1.0)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, bold=header or label, colour=WHITE if header else BLACK)
    set_cell_fill(cell, "1F4E78" if header else ("F2F7FB" if shaded else "FFFFFF"))


def add_comparison_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=4)
    for cell, text in zip(table.rows[0].cells, TABLE_HEADERS, strict=True):
        fill_cell(cell, text, header=True)
    set_repeat_table_header(table.rows[0])

    for row_index, values in enumerate(TABLE_ROWS, start=1):
        row = table.add_row()
        for column_index, (cell, text) in enumerate(zip(row.cells, values, strict=True)):
            fill_cell(
                cell,
                text,
                label=column_index == 0,
                shaded=row_index % 2 == 0,
            )
    set_table_geometry(table)


def build_donor(path: Path) -> None:
    document = Document(REFERENCE)
    clear_document_body(document)

    first = document.sections[0]
    first.orientation = WD_ORIENT.PORTRAIT
    first.page_width = Mm(210)
    first.page_height = Mm(297)
    first.top_margin = Inches(1)
    first.bottom_margin = Inches(1)
    first.left_margin = Inches(1)
    first.right_margin = Inches(1)

    add_title(document, "Weekly Update 31 August to 4 September")
    add_intro(
        document,
        "The main changes added multilingual narration to Fundal Reflex, improved video and subtitle controls and reduced scheduled reviews to six runs per week.",
    )
    for heading, items in SECTIONS:
        add_heading(document, heading)
        for item in items:
            add_bullet(document, item)

    landscape = document.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width = Mm(297)
    landscape.page_height = Mm(210)
    landscape.top_margin = Inches(0.55)
    landscape.bottom_margin = Inches(0.55)
    landscape.left_margin = Inches(0.55)
    landscape.right_margin = Inches(0.55)
    landscape.header_distance = Inches(0.3)
    landscape.footer_distance = Inches(0.3)

    add_heading(document, "Scheduled Tasks Comparison")
    add_comparison_table(document)

    document.core_properties.title = "Weekly Update 31 August to 4 September"
    document.save(path)


def update_core_xml(source: bytes) -> bytes:
    parser = etree.XMLParser(remove_blank_text=False)
    root = etree.fromstring(source, parser)
    title = root.find("dc:title", NS)
    if title is not None:
        title.text = "Weekly Update 31 August to 4 September"
    modified = root.find("dcterms:modified", NS)
    if modified is not None:
        modified.text = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")


def build() -> None:
    if sha256(REFERENCE) != EXPECTED_SHA256:
        raise RuntimeError("Reference SHA-256 no longer matches the distilled template")

    TASK_DIR.mkdir(parents=True, exist_ok=True)
    donor_handle, donor_name = tempfile.mkstemp(prefix="donor-", suffix=".docx", dir=TASK_DIR)
    os.close(donor_handle)
    donor = Path(donor_name)
    output_handle, output_name = tempfile.mkstemp(prefix="changelog-", suffix=".docx", dir=OUTPUT.parent)
    os.close(output_handle)
    staged_output = Path(output_name)

    try:
        build_donor(donor)
        with zipfile.ZipFile(donor, "r") as donor_zip:
            document_xml = donor_zip.read("word/document.xml")
        with zipfile.ZipFile(REFERENCE, "r") as source_zip, zipfile.ZipFile(staged_output, "w") as output_zip:
            for info in source_zip.infolist():
                data = source_zip.read(info.filename)
                if info.filename == "word/document.xml":
                    data = document_xml
                elif info.filename == "docProps/core.xml":
                    data = update_core_xml(data)
                output_zip.writestr(info, data)
        shutil.move(str(staged_output), OUTPUT)
    finally:
        donor.unlink(missing_ok=True)
        staged_output.unlink(missing_ok=True)

    if sha256(REFERENCE) != EXPECTED_SHA256:
        raise RuntimeError("Reference changed during authoring")
    print(OUTPUT)


if __name__ == "__main__":
    build()
