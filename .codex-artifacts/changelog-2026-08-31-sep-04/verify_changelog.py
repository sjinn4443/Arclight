from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn


REFERENCE = Path(r"C:\Users\young\Downloads\15 - 19 June.docx")
OUTPUT = Path(r"D:\Arclight\Arclight_App_NewApp_ChatRefact\Arclight_App\31 August - 4 September.docx")
EXPECTED_REFERENCE_SHA256 = "0B19F0022172001957D1D9703465646C89E02A3AECA2DB284FFAE95386EC5122"

EXPECTED_HEADERS = ["Area", "Before", "Now", "Practical benefit"]
EXPECTED_ROWS = [
    ["Total scheduled runs", "About 20 runs per week", "Six runs per week", "Uses less allowance and focuses on essential checks"],
    ["Operations monitoring", "Daily at 04:00 and weekdays at 08:30", "Monday, Wednesday and Friday at 08:30", "Removes repetitive status checks"],
    ["Detecting new updates", "Mainly checked recent activity on main", "Checks changes since the previous run and distinguishes pushes to other branches", "Less likely to miss recently pushed work"],
    ["Product and UI review", "Wednesday and translation checks ran separately", "Combined into one Wednesday 04:00 review", "UI, accessibility, screen sizes and language findings arrive together"],
    ["Full translation review", "Split across four monthly runs", "Included in the first Wednesday report each month", "Retains full monthly coverage with fewer runs"],
    ["Security and housekeeping", "Security ran on Friday while repository hygiene ran on Monday", "Combined into Friday's Security + Hygiene review", "Security, refactoring and unnecessary-file candidates appear in one report"],
    ["Usage control", "Work could stop when a 10% monthly allowance could not be calculated", "Scope adjusts gradually using actual rolling seven-day usage", "Prevents jobs being skipped despite available capacity"],
    ["Report format", "Most results were brief chat messages", "Major reviews provide detailed chat reports and a DOCX when supported", "Results are easier to understand and share"],
    ["PostgreSQL DOCX", "Could not be created without the reporting connection", "Ready for private server-side generation once connected", "It will remain unavailable until Arclight Reporting is connected"],
    ["Safety", "Read-only guidance existed", "Explicit read-only restrictions now apply to every active job", "Scheduled work cannot change code, deployments or database records"],
    ["Agent names", "Run days were not obvious", "[Mon/Wed/Fri]Operations, [Wed]Product + Language, [Fri]Security + Hygiene, [Fri]Weekly Reports", "Each name clearly shows when the agent runs and what it does"],
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


document = Document(OUTPUT)
assert document.paragraphs[0].text == "Weekly Update 31 August to 4 September"
assert document.paragraphs[0].style.name == "Title"
assert len(document.sections) == 2
assert document.sections[0].orientation == WD_ORIENT.PORTRAIT
assert document.sections[1].orientation == WD_ORIENT.LANDSCAPE
assert round(document.sections[0].left_margin.inches, 2) == 1.00
assert round(document.sections[1].left_margin.inches, 2) == 0.55

body_items = [p for p in document.paragraphs if p.style.name == "List Paragraph"]
assert len(body_items) == 12
assert {p._p.pPr.numPr.numId.val for p in body_items} == {36}
assert any("BIO" in p.text and "in progress" in p.text for p in body_items)

assert len(document.tables) == 1
table = document.tables[0]
assert len(table.rows) == 12
assert len(table.columns) == 4
assert [cell.text for cell in table.rows[0].cells] == EXPECTED_HEADERS
assert [[cell.text for cell in row.cells] for row in table.rows[1:]] == EXPECTED_ROWS

tbl_pr = table._tbl.tblPr
tbl_w = tbl_pr.find(qn("w:tblW"))
assert tbl_w is not None and tbl_w.get(qn("w:w")) == "15240" and tbl_w.get(qn("w:type")) == "dxa"
assert [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid.gridCol_lst] == [2160, 3420, 4800, 4860]
assert table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None

for row in table.rows:
    assert row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
    assert row._tr.get(qn("w:h")) is None
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        assert tc_pr.find(qn("w:tcMar")) is not None
        assert tc_pr.find(qn("w:tcW")) is not None

all_text = [p.text for p in document.paragraphs]
all_text.extend(cell.text for row in table.rows for cell in row.cells)
assert not [text for text in all_text if re.search(r",\s+(and|or)\s+", text, re.I)]
assert not any("PLACEHOLDER" in text or "TODO" in text for text in all_text)

with zipfile.ZipFile(REFERENCE) as reference_zip, zipfile.ZipFile(OUTPUT) as output_zip:
    reference_parts = set(reference_zip.namelist())
    output_parts = set(output_zip.namelist())
    assert reference_parts == output_parts
    changed_parts = [
        name
        for name in sorted(reference_parts)
        if hashlib.sha256(reference_zip.read(name)).digest()
        != hashlib.sha256(output_zip.read(name)).digest()
    ]
    assert changed_parts == ["docProps/core.xml", "word/document.xml"]

assert sha256(REFERENCE) == EXPECTED_REFERENCE_SHA256
print("paragraphs", len(document.paragraphs))
print("body_list_items", len(body_items))
print("sections", len(document.sections))
print("table_rows", len(table.rows))
print("table_columns", len(table.columns))
print("changed_parts", changed_parts)
print("output_sha256", sha256(OUTPUT))
