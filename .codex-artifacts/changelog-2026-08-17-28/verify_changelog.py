from __future__ import annotations

import hashlib
import os
import re
import zipfile
from pathlib import Path

from docx import Document


REFERENCE = Path(r"C:\Users\young\Downloads\15 - 19 June.docx")
OUTPUT = Path(r"D:\Arclight\Arclight_App_NewApp_ChatRefact\Arclight_App\17 - 28 August.docx")


document = Document(OUTPUT)
items = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "List Paragraph"]

print("paragraphs", len(document.paragraphs))
print("tables", len(document.tables))
print("sections", len(document.sections))
print("title", repr(document.paragraphs[0].text))
print("list_items", len(items))
print("num_ids", sorted({paragraph._p.pPr.numPr.numId.val for paragraph in items}))
print("empty_paragraphs", sum(not paragraph.text for paragraph in document.paragraphs))
print(
    "oxford_comma_candidates",
    [paragraph.text for paragraph in document.paragraphs if re.search(r",\s+(and|or)\s+", paragraph.text, re.I)],
)
print("TEXT")
for index, paragraph in enumerate(document.paragraphs):
    print(f"{index:02d}|{paragraph.style.name}|{paragraph.text}")

with zipfile.ZipFile(REFERENCE) as reference_zip, zipfile.ZipFile(OUTPUT) as output_zip:
    reference_parts = set(reference_zip.namelist())
    output_parts = set(output_zip.namelist())
    changed_parts = [
        name
        for name in sorted(reference_parts & output_parts)
        if hashlib.sha256(reference_zip.read(name)).digest()
        != hashlib.sha256(output_zip.read(name)).digest()
    ]
    print("same_parts", reference_parts == output_parts)
    print("part_count", len(output_parts))
    print("changed_parts", changed_parts)

print("output_size", os.path.getsize(OUTPUT))
